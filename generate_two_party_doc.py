import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = Document()

# Set clean compact margins for single page layout
for section in doc.sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

# Header Section
header_p = doc.add_paragraph()
header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
header_p.paragraph_format.space_after = Pt(2)
header_p.paragraph_format.space_before = Pt(0)
header_p.paragraph_format.line_spacing = 1.05

r_title = header_p.add_run('PARROW SKILLS\n')
r_title.font.name = 'Arial'
r_title.font.size = Pt(17)
r_title.font.bold = True
r_title.font.color.rgb = RGBColor(16, 44, 87)

r_sub = header_p.add_run(
    '1-1386-B, BC Colony, Pamidi, Ananthapur, Andhra Pradesh - 515775, India\n'
    'Email: tameemansarkhan@gmail.com  |  Mobile: +91 8008081298\n'
    'MSME / Udyam Registration No: UDYAM-AP-01-0034849'
)
r_sub.font.name = 'Arial'
r_sub.font.size = Pt(8.5)
r_sub.font.color.rgb = RGBColor(90, 90, 90)

# Divider line
p_div = doc.add_paragraph()
p_div.paragraph_format.space_before = Pt(2)
p_div.paragraph_format.space_after = Pt(6)
p_div_border = OxmlElement('w:pBdr')
bottom_border = OxmlElement('w:bottom')
bottom_border.set(qn('w:val'), 'single')
bottom_border.set(qn('w:sz'), '10')
bottom_border.set(qn('w:space'), '1')
bottom_border.set(qn('w:color'), '102C57')
p_div_border.append(bottom_border)
p_div._p.get_or_add_pPr().append(p_div_border)

# Document Title
p_doc_title = doc.add_paragraph()
p_doc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_doc_title.paragraph_format.space_after = Pt(6)
p_doc_title.paragraph_format.space_before = Pt(0)
r_doctitle = p_doc_title.add_run('MUTUAL BRAND AUTHORIZATION & DISTRIBUTION AGREEMENT')
r_doctitle.font.name = 'Arial'
r_doctitle.font.size = Pt(11)
r_doctitle.font.bold = True
r_doctitle.font.underline = True
r_doctitle.font.color.rgb = RGBColor(16, 44, 87)

# Date & Reference
p_to = doc.add_paragraph()
p_to.paragraph_format.space_after = Pt(4)
p_to.paragraph_format.space_before = Pt(0)
p_to.paragraph_format.line_spacing = 1.05
r_date = p_to.add_run('Effective Date: 29th August 2026   |   Appeal Case Ref: [5-6786000041693]\n')
r_date.font.bold = True
r_date.font.size = Pt(9)
r_to = p_to.add_run('To: Google Play App Review & Policy Team, Google LLC')
r_to.font.name = 'Arial'
r_to.font.size = Pt(9)

# Parties definition
p_parties = doc.add_paragraph()
p_parties.paragraph_format.space_after = Pt(4)
p_parties.paragraph_format.space_before = Pt(0)
p_parties.paragraph_format.line_spacing = 1.05
p_parties.add_run('This Agreement is entered into between:\n').font.name = 'Arial'
r_p1 = p_parties.add_run('1. FIRST PARTY (Brand & Content Owner): ')
r_p1.font.bold = True
r_p1.font.size = Pt(8.5)
p_parties.add_run('PARROW SKILLS (Udyam Reg. No: UDYAM-AP-01-0034849), represented by its Proprietor, Tameem Ansar Khan, having its registered office at 1-1386-B, BC Colony, Pamidi, Ananthapur, Andhra Pradesh - 515775.\n').font.size = Pt(8.5)
r_p2 = p_parties.add_run('2. SECOND PARTY (Authorized Developer): ')
r_p2.font.bold = True
r_p2.font.size = Pt(8.5)
p_parties.add_run('Premium Business (CODTECH IT SOLUTIONS), Google Play Developer Email: codtechitsolutions@gmail.com.\n').font.size = Pt(8.5)

# Agreement Clauses
bullet_points = [
    ('1. Grant of License & Authorization: ', 'The First Party hereby grants the Second Party the full legal authorization and license to develop, upload, manage, publish, and distribute the official mobile application titled "ParrowSkills" (Package ID: com.parrowskills.app) on the Google Play Store.'),
    ('2. Use of Brand Assets & Logos: ', 'The First Party confirms that the Second Party is fully authorized to use all associated brand names, trademarks, logos, icons, descriptions, and media assets belonging to "ParrowSkills" on Google Play.'),
    ('3. Mutual Consent & Acceptance: ', 'Both Parties mutually agree and confirm that this publishing relationship is legitimate, legally binding, and authorized by both entities.')
]

for title, desc in bullet_points:
    p_bp = doc.add_paragraph()
    p_bp.paragraph_format.left_indent = Inches(0.15)
    p_bp.paragraph_format.space_after = Pt(2)
    p_bp.paragraph_format.space_before = Pt(0)
    p_bp.paragraph_format.line_spacing = 1.05
    r_t = p_bp.add_run(title)
    r_t.font.bold = True
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(8.5)
    r_d = p_bp.add_run(desc)
    r_d.font.name = 'Arial'
    r_d.font.size = Pt(8.5)

# Details Table
table = doc.add_table(rows=4, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table_data = [
    ('App Title & Package Name', 'ParrowSkills (com.parrowskills.app)'),
    ('Brand Owner Entity (Licensor)', 'PARROW SKILLS (MSME: UDYAM-AP-01-0034849)'),
    ('Authorized Developer Account', 'Premium Business'),
    ('Developer Account Email', 'codtechitsolutions@gmail.com')
]

for row_idx, (k, v) in enumerate(table_data):
    row = table.rows[row_idx]
    cell_0 = row.cells[0]
    cell_1 = row.cells[1]
    cell_0.width = Inches(2.6)
    cell_1.width = Inches(4.3)
    p0 = cell_0.paragraphs[0]
    p0.paragraph_format.space_after = Pt(1)
    p0.paragraph_format.space_before = Pt(1)
    r0 = p0.add_run(k)
    r0.font.bold = True
    r0.font.size = Pt(8)
    
    p1 = cell_1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(1)
    p1.paragraph_format.space_before = Pt(1)
    r1 = p1.add_run(v)
    r1.font.bold = True
    r1.font.size = Pt(8)
    r1.font.color.rgb = RGBColor(16, 44, 87)

for row in table.rows:
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')
        for border_name in ['top', 'left', 'bottom', 'right']:
            b = OxmlElement(f'w:{border_name}')
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), '4')
            b.set(qn('w:color'), 'D0D0D0')
            tcBorders.append(b)
        tcPr.append(tcBorders)

# Enclosure Note
p_attach = doc.add_paragraph()
p_attach.paragraph_format.space_before = Pt(3)
p_attach.paragraph_format.space_after = Pt(4)
r_att = p_attach.add_run('Enclosures: 1. MSME / Udyam Certificate (UDYAM-AP-01-0034849) | 2. Government Photo ID of Brand Owner.')
r_att.font.italic = True
r_att.font.size = Pt(8)
r_att.font.color.rgb = RGBColor(80, 80, 80)

# TWO-PARTY SIGNATURE SECTION TABLE
p_sig_heading = doc.add_paragraph()
p_sig_heading.paragraph_format.space_before = Pt(2)
p_sig_heading.paragraph_format.space_after = Pt(2)
r_sh = p_sig_heading.add_run('IN WITNESS WHEREOF, BOTH PARTIES HAVE EXECUTED THIS AGREEMENT:')
r_sh.font.bold = True
r_sh.font.size = Pt(8.5)
r_sh.font.color.rgb = RGBColor(16, 44, 87)

sig_table = doc.add_table(rows=1, cols=2)
sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
row_sig = sig_table.rows[0]

# Left: Party 1 (Brand Owner)
cell_p1 = row_sig.cells[0]
cell_p1.width = Inches(3.45)
p_p1 = cell_p1.paragraphs[0]
p_p1.paragraph_format.line_spacing = 1.05
p_p1.paragraph_format.space_after = Pt(0)
p_p1.paragraph_format.space_before = Pt(0)
p_p1.add_run('FIRST PARTY (Brand Owner):\n').font.bold = True
p_p1.add_run('For PARROW SKILLS\n\n\n')
p_p1.add_run('___________________________________\n')
p_p1.add_run('Signature: Tameem Ansar Khan\n').font.bold = True
p_p1.add_run('Title: Proprietor / Brand Owner\nMobile: +91 8008081298\nEmail: tameemansarkhan@gmail.com')
for run in p_p1.runs:
    if run.font.size is None:
        run.font.size = Pt(8)

# Right: Party 2 (Developer)
cell_p2 = row_sig.cells[1]
cell_p2.width = Inches(3.45)
p_p2 = cell_p2.paragraphs[0]
p_p2.paragraph_format.line_spacing = 1.05
p_p2.paragraph_format.space_after = Pt(0)
p_p2.paragraph_format.space_before = Pt(0)
p_p2.add_run('SECOND PARTY (Developer):\n').font.bold = True
p_p2.add_run('For Premium Business\n\n\n')
p_p2.add_run('___________________________________\n')
p_p2.add_run('Signature: Authorized Signatory\n').font.bold = True
p_p2.add_run('Title: Lead Developer / Account Owner\nDeveloper Email: codtechitsolutions@gmail.com')
for run in p_p2.runs:
    if run.font.size is None:
        run.font.size = Pt(8)

# Remove border around signature table
for cell in row_sig.cells:
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        b = OxmlElement(f'w:{border_name}')
        b.set(qn('w:val'), 'none')
        tcBorders.append(b)
    tcPr.append(tcBorders)

out_path = r'd:\Company Projects\Easy Labour Finder\ParrowSkills_Two_Party_Agreement.docx'
doc.save(out_path)
print("Two-party agreement successfully created at:", out_path)

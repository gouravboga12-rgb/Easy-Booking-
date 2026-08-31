import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = Document()

# Set compact margins to guarantee single page layout
for section in doc.sections:
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

# Header Section (Company Letterhead Header)
header_p = doc.add_paragraph()
header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
header_p.paragraph_format.space_after = Pt(2)
header_p.paragraph_format.space_before = Pt(0)
header_p.paragraph_format.line_spacing = 1.05

r_title = header_p.add_run('PARROW SKILLS\n')
r_title.font.name = 'Arial'
r_title.font.size = Pt(18)
r_title.font.bold = True
r_title.font.color.rgb = RGBColor(16, 44, 87)

r_sub = header_p.add_run(
    '1-1386-B, BC Colony, Pamidi, Ananthapur, Andhra Pradesh - 515775, India\n'
    'Email: tameemansarkhan@gmail.com  |  Mobile: +91 8008081298\n'
    'MSME / Udyam Registration No: UDYAM-AP-01-0034849'
)
r_sub.font.name = 'Arial'
r_sub.font.size = Pt(8.5)
r_sub.font.color.rgb = RGBColor(90, 90, 90)

# Divider line
p_div = doc.add_paragraph()
p_div.paragraph_format.space_before = Pt(2)
p_div.paragraph_format.space_after = Pt(8)
p_div_border = OxmlElement('w:pBdr')
bottom_border = OxmlElement('w:bottom')
bottom_border.set(qn('w:val'), 'single')
bottom_border.set(qn('w:sz'), '10')
bottom_border.set(qn('w:space'), '1')
bottom_border.set(qn('w:color'), '102C57')
p_div_border.append(bottom_border)
p_div._p.get_or_add_pPr().append(p_div_border)

# Document Title
p_doc_title = doc.add_paragraph()
p_doc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_doc_title.paragraph_format.space_after = Pt(8)
p_doc_title.paragraph_format.space_before = Pt(0)
r_doctitle = p_doc_title.add_run('OFFICIAL BRAND AUTHORIZATION & DECLARATION LETTER')
r_doctitle.font.name = 'Arial'
r_doctitle.font.size = Pt(11.5)
r_doctitle.font.bold = True
r_doctitle.font.underline = True
r_doctitle.font.color.rgb = RGBColor(16, 44, 87)

# Date & Recipient
p_to = doc.add_paragraph()
p_to.paragraph_format.space_after = Pt(6)
p_to.paragraph_format.space_before = Pt(0)
p_to.paragraph_format.line_spacing = 1.05
r_date = p_to.add_run('Date: 29th August 2026\n')
r_date.font.bold = True
r_date.font.size = Pt(9.5)
r_to = p_to.add_run(
    'To: The Google Play App Review & Policy Team, Google LLC'
)
r_to.font.name = 'Arial'
r_to.font.size = Pt(9.5)

# Subject
p_sub = doc.add_paragraph()
p_sub.paragraph_format.space_after = Pt(6)
p_sub.paragraph_format.space_before = Pt(0)
p_sub.paragraph_format.line_spacing = 1.05
r_sub_label = p_sub.add_run('Subject: ')
r_sub_label.font.bold = True
r_sub_label.font.size = Pt(9.5)
r_sub_text = p_sub.add_run(
    'Brand Authorization & Ownership Verification for Mobile App "ParrowSkills" (Package ID: com.parrowskills.app)'
)
r_sub_text.font.bold = True
r_sub_text.font.size = Pt(9.5)

# Body Paragraph 1
p_b1 = doc.add_paragraph()
p_b1.paragraph_format.space_after = Pt(5)
p_b1.paragraph_format.space_before = Pt(0)
p_b1.paragraph_format.line_spacing = 1.05
p_b1.add_run('Dear Google Play Team,\n').font.name = 'Arial'
p_b1.add_run('We, ').font.name = 'Arial'
r_bname = p_b1.add_run('PARROW SKILLS')
r_bname.font.bold = True
p_b1.add_run(
    ', an officially registered business enterprise under the Ministry of Micro, Small & Medium Enterprises (MSME), Government of India (Udyam Reg. No: '
)
r_reg = p_b1.add_run('UDYAM-AP-01-0034849')
r_reg.font.bold = True
p_b1.add_run(
    '), having our registered address at 1-1386-B, BC Colony, Pamidi, Ananthapur, Andhra Pradesh - 515775, India, do hereby solemnly declare and confirm that:'
)

# Key clauses
bullet_points = [
    ('1. Legal Brand Ownership: ', 'PARROW SKILLS is the legitimate creator, rightful owner, and legal proprietor of the brand name, trade identity, graphics, and official logo of "ParrowSkills".'),
    ('2. Developer Authorization: ', 'We have officially authorized the Google Play Developer Account "Premium Business" (Email: codtechitsolutions@gmail.com) to develop, publish, manage, update, and distribute our official Android mobile application titled "ParrowSkills" (Package Name: com.parrowskills.app) on the Google Play Store on our behalf.'),
    ('3. Intellectual Property Rights: ', 'The authorized developer account is granted full authorization to utilize all associated brand names, trademarks, logos, icons, descriptions, and promotional assets belonging to PARROW SKILLS for the purpose of operating this application.'),
    ('4. Validity & Confirmation: ', 'This authorization is active, legally binding, granted with our full consent, and is fully recognized by our enterprise.')
]

for title, desc in bullet_points:
    p_bp = doc.add_paragraph()
    p_bp.paragraph_format.left_indent = Inches(0.15)
    p_bp.paragraph_format.space_after = Pt(3)
    p_bp.paragraph_format.space_before = Pt(0)
    p_bp.paragraph_format.line_spacing = 1.05
    r_t = p_bp.add_run(title)
    r_t.font.bold = True
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(9)
    r_d = p_bp.add_run(desc)
    r_d.font.name = 'Arial'
    r_d.font.size = Pt(9)

# Details Table
p_table_space = doc.add_paragraph()
p_table_space.paragraph_format.space_before = Pt(2)
p_table_space.paragraph_format.space_after = Pt(2)

table = doc.add_table(rows=4, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table_data = [
    ('Application Title', 'ParrowSkills'),
    ('Application Package Name (ID)', 'com.parrowskills.app'),
    ('Authorized Developer Account Name', 'Premium Business'),
    ('Authorized Developer Account Email', 'codtechitsolutions@gmail.com')
]

for row_idx, (k, v) in enumerate(table_data):
    row = table.rows[row_idx]
    cell_0 = row.cells[0]
    cell_1 = row.cells[1]
    cell_0.width = Inches(2.7)
    cell_1.width = Inches(4.2)
    
    p0 = cell_0.paragraphs[0]
    p0.paragraph_format.space_after = Pt(2)
    p0.paragraph_format.space_before = Pt(2)
    r0 = p0.add_run(k)
    r0.font.bold = True
    r0.font.size = Pt(8.5)
    
    p1 = cell_1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.space_before = Pt(2)
    r1 = p1.add_run(v)
    r1.font.bold = True
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = RGBColor(16, 44, 87)

# Set clean subtle border
for row in table.rows:
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')
        for border_name in ['top', 'left', 'bottom', 'right']:
            b = OxmlElement(f'w:{border_name}')
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), '4')
            b.set(qn('w:color'), 'D0D0D0')
            tcBorders.append(b)
        tcPr.append(tcBorders)

# Enclosure note
p_attach = doc.add_paragraph()
p_attach.paragraph_format.space_before = Pt(4)
p_attach.paragraph_format.space_after = Pt(6)
r_att = p_attach.add_run('Enclosure: Official Government MSME / Udyam Registration Certificate (UDYAM-AP-01-0034849).')
r_att.font.italic = True
r_att.font.size = Pt(8.5)
r_att.font.color.rgb = RGBColor(90, 90, 90)

# Signatory block
p_sig = doc.add_paragraph()
p_sig.paragraph_format.space_after = Pt(1)
p_sig.paragraph_format.space_before = Pt(0)
p_sig.paragraph_format.line_spacing = 1.05
r_sin = p_sig.add_run('Sincerely,\nFor ')
r_sin.font.size = Pt(9.5)
r_comp = p_sig.add_run('PARROW SKILLS\n\n\n')
r_comp.font.bold = True
r_comp.font.size = Pt(9.5)

p_sign_line = doc.add_paragraph()
p_sign_line.paragraph_format.space_after = Pt(2)
p_sign_line.paragraph_format.space_before = Pt(0)
p_sign_line.paragraph_format.line_spacing = 1.05
r_line = p_sign_line.add_run('_________________________________________\n(Signature of Authorized Signatory)')
r_line.font.size = Pt(8.5)
r_line.font.color.rgb = RGBColor(90, 90, 90)

p_signatory = doc.add_paragraph()
p_signatory.paragraph_format.line_spacing = 1.05
p_signatory.paragraph_format.space_after = Pt(0)
p_signatory.paragraph_format.space_before = Pt(0)
r_name = p_signatory.add_run('Name: Tameem Ansar Khan\n')
r_name.font.bold = True
r_name.font.size = Pt(9.5)
r_desig = p_signatory.add_run(
    'Designation: Proprietor / Authorized Representative\n'
    'Enterprise: PARROW SKILLS\n'
    'Email: tameemansarkhan@gmail.com  |  Mobile: +91 8008081298'
)
r_desig.font.size = Pt(8.5)

output_path_updated = r'd:\Company Projects\Easy Labour Finder\ParrowSkills_Brand_Authorization_Letter_Updated.docx'
doc.save(output_path_updated)
print("Saved to:", output_path_updated)

try:
    output_path = r'd:\Company Projects\Easy Labour Finder\ParrowSkills_Brand_Authorization_Letter.docx'
    doc.save(output_path)
    print("Also overwritten:", output_path)
except Exception as e:
    print("Could not overwrite open file, please use Updated file.")
